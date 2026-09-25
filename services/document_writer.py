import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor


class DocumentWriter:

    def create_document(
        self, title: str, sections: list[dict], sources: list[dict], output_file: str
    ) -> str:
        doc = Document()

        # Document Title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = 1  # Centered

        # Add Sections
        for section in sections:
            doc.add_heading(section["heading"], level=1)
            doc.add_paragraph(section["content"])

        # Add Sources / References Section if sources exist
        if sources:
            doc.add_heading("References & Sources Cited", level=1)
            # Remove duplicate URLs while preserving order
            unique_sources = []
            seen_urls = set()
            for src in sources:
                if src["url"] not in seen_urls:
                    seen_urls.add(src["url"])
                    unique_sources.append(src)

            for src in unique_sources:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{src['title']}\n").bold = True
                p.add_run(f"Source URL: {src['url']}")

        os.makedirs(os.path.dirname(output_file), exist_ok=True)
        doc.save(output_file)
        return output_file