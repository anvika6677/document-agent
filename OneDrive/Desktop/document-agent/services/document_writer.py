from docx import Document
from docx.shared import Pt


class DocumentWriter:

    def __init__(self):
        self.document = Document()

    def write(self, title: str, sections: list, output_file: str):

        # Document Title
        heading = self.document.add_heading(title, level=1)
        heading.runs[0].font.size = Pt(20)

        self.document.add_paragraph()

        # Add every generated section
        for section in sections:

            self.document.add_heading(section["heading"], level=2)

            self.document.add_paragraph(section["content"])

            self.document.add_paragraph()

        self.document.save(output_file)

        return output_file